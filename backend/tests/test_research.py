"""Fase 5: episodios automáticos y manuales, timeline, comparación, memos, entrevistas, exportación."""

from __future__ import annotations

import csv
import io
import json

import httpx
import pytest
from app.modules.learning.models import Task
from app.modules.ops.models import AuditLog
from app.seed.catalog import seed_catalog
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from tests.conftest import LoginFn, SeededWorld


@pytest.fixture
async def catalog(db: AsyncSession, world: SeededWorld) -> dict[str, str]:
    await seed_catalog(db)
    await db.commit()
    return {t.code: str(t.id) for t in (await db.execute(select(Task))).scalars()}


async def build_episode(
    client: httpx.AsyncClient, headers: dict[str, str], task_id: str, code: str = "B-FIG-01"
) -> str:
    session = (await client.post("/api/v1/sessions", json={"task_codes": [code]}, headers=headers)).json()
    sid = session["id"]
    await client.post(
        f"/api/v1/sessions/{sid}/events", json={"event_type": "TASK_OPENED", "task_id": task_id}, headers=headers
    )
    await client.post(
        f"/api/v1/sessions/{sid}/responses",
        json={
            "task_id": task_id,
            "question_id": "q3",
            "representation": "NUMERIC",
            "content": {"value": 41},
            "client_meta": {"time_on_task_ms": 200000},
        },
        headers=headers,
    )
    offer = (
        await client.post(
            f"/api/v1/sessions/{sid}/help", json={"task_id": task_id, "question_id": "q3"}, headers=headers
        )
    ).json()
    await client.post(
        f"/api/v1/sessions/{sid}/scaffold-events/{offer['scaffold_event_id']}/feedback",
        json={"action": "ACCEPTED"},
        headers=headers,
    )
    await client.post(
        f"/api/v1/sessions/{sid}/events",
        json={"event_type": "REPRESENTATION_CHANGED", "task_id": task_id, "representation": "TABULAR"},
        headers=headers,
    )
    await client.post(
        f"/api/v1/sessions/{sid}/responses",
        json={
            "task_id": task_id,
            "question_id": "q3",
            "representation": "NUMERIC",
            "content": {"value": 38},
            "client_meta": {"time_on_task_ms": 200000},
        },
        headers=headers,
    )
    return sid


async def test_automatic_episode_reconstructs_trajectory(
    client: httpx.AsyncClient, world: SeededWorld, login: LoginFn, catalog: dict[str, str]
) -> None:
    student = await login(world.student_username)
    sid = await build_episode(client, student, catalog["B-FIG-01"])
    researcher = await login(world.researcher_email)

    episodes = (await client.get("/api/v1/research/episodes", params={"session_id": sid}, headers=researcher)).json()
    assert len(episodes) == 1
    ep = episodes[0]
    assert ep["trigger"] == "HELP_REQUEST"
    assert ep["status"] == "CLOSED" and ep["close_reason"] == "RESOLVED"
    assert ep["participant_code"] == world.student_participant_code
    summary = ep["summary"]
    assert summary["system_generated"] is True
    assert summary["help_accepted"] == 1 and summary["responses"] == 2
    assert summary["responses_evaluated_correct"] == 1 and summary["responses_evaluated_incorrect"] == 1

    detail = (await client.get(f"/api/v1/research/episodes/{ep['id']}", headers=researcher)).json()
    phases = [s["phase"] for s in detail["timeline"]]
    # Dificultad → Solicitud de ayuda → Ayuda → Reacción → Cambio → Nueva respuesta.
    assert phases == [
        "TAREA",
        "RESPUESTA",
        "DIFICULTAD",
        "SOLICITUD DE AYUDA",
        "AYUDA",
        "REACCIÓN",
        "CAMBIO",
        "NUEVA RESPUESTA",
    ]
    actors = {s["event_type"]: s["actor"] for s in detail["timeline"]}
    assert actors["HELP_OFFERED"] == "SYSTEM" and actors["STUDENT_RESPONSE"] == "STUDENT"
    help_step = next(s for s in detail["timeline"] if s["event_type"] == "HELP_OFFERED")
    assert help_step["scaffold_decision"]["reason"]
    request_step = next(s for s in detail["timeline"] if s["event_type"] == "HELP_REQUESTED")
    assert request_step["system_interpretation"]["label"] == "interpretación operativa del sistema"
    assert detail["context"]["task"]["code"] == "B-FIG-01"
    # Sin datos personales.
    assert "stu1" not in json.dumps(detail) and "@" not in json.dumps(detail)

    teacher = await login(world.teacher_email)
    assert (await client.get(f"/api/v1/research/episodes/{ep['id']}", headers=teacher)).status_code == 200
    assert (await client.get("/api/v1/research/episodes", headers=student)).status_code == 403


async def test_manual_episode_compare_memos_interviews(
    client: httpx.AsyncClient, world: SeededWorld, login: LoginFn, catalog: dict[str, str]
) -> None:
    student = await login(world.student_username)
    sid = await build_episode(client, student, catalog["B-FIG-01"])
    await client.patch(f"/api/v1/sessions/{sid}/end", json={"status": "COMPLETED"}, headers=student)
    sid2 = await build_episode(client, student, catalog["A-NUM-01"], code="A-NUM-01")
    researcher = await login(world.researcher_email)

    manual = await client.post(
        "/api/v1/research/episodes",
        json={"session_id": sid2, "task_id": catalog["A-NUM-01"], "from_sequence": 2, "to_sequence": 4},
        headers=researcher,
    )
    assert manual.status_code == 201, manual.text
    assert manual.json()["trigger"] == "MANUAL" and manual.json()["interaction_count"] == 3

    eps = (await client.get("/api/v1/research/episodes", headers=researcher)).json()
    auto = [e for e in eps if e["trigger"] != "MANUAL"]
    comparison = await client.get(
        "/api/v1/research/episodes/compare", params={"a": auto[0]["id"], "b": auto[1]["id"]}, headers=researcher
    )
    assert comparison.status_code == 200
    body = comparison.json()
    assert set(body["dimensions"]) == {
        "contexto",
        "estrategia",
        "representación",
        "ayuda",
        "reacción",
        "intervención_docente",
        "trayectoria",
        "cierre",
    }
    assert "no produce conclusiones" in body["note"]

    memo = await client.post(
        "/api/v1/research/memos",
        json={
            "title": "Cambio a tabla tras la pista",
            "episode_id": auto[0]["id"],
            "observation": "Tras aceptar la pista cambia a tabla.",
            "interpretation": "Posible uso de la tabla como apoyo.",
            "emerging_question": "¿La tabla media la relación funcional?",
            "negative_case": "En el episodio B no cambia de representación.",
            "possible_category": "Mediación representacional",
            "tags": ["tabla"],
        },
        headers=researcher,
    )
    assert memo.status_code == 201, memo.text
    assert memo.json()["participant_code"] == world.student_participant_code
    assert memo.json()["source"] == "RESEARCHER"
    detail = (await client.get(f"/api/v1/research/episodes/{auto[0]['id']}", headers=researcher)).json()
    assert detail["memos"][0]["possible_category"] == "Mediación representacional"

    updated = await client.put(
        f"/api/v1/research/memos/{memo.json()['id']}", json={"title": "Editado", "tags": []}, headers=researcher
    )
    assert updated.status_code == 200 and updated.json()["title"] == "Editado"
    admin = await login(world.admin_email)
    assert (await client.post("/api/v1/research/memos", json={"title": "Memo admin"}, headers=admin)).status_code == 403
    assert (
        await client.post(
            "/api/v1/research/memos", json={"title": "Memo docente"}, headers=await login(world.teacher_email)
        )
    ).status_code == 403

    interview = await client.post(
        "/api/v1/research/interviews",
        json={
            "title": "Entrevista posterior",
            "interviewee_kind": "STUDENT",
            "participant_code": world.student_participant_code,
            "episode_id": auto[0]["id"],
            "conducted_at": "2026-09-28T10:00:00Z",
            "responses": [
                {"question": "¿Qué te ayudó?", "answer": "Hacer la tabla", "related_episode_id": auto[0]["id"]}
            ],
        },
        headers=researcher,
    )
    assert interview.status_code == 201, interview.text
    added = await client.post(
        f"/api/v1/research/interviews/{interview.json()['id']}/responses",
        json={"question": "¿Y después?", "answer": "Vi que suma 3"},
        headers=researcher,
    )
    assert added.json()["order_index"] == 2
    listed = (await client.get("/api/v1/research/interviews", headers=researcher)).json()
    assert listed[0]["participant_code"] == world.student_participant_code and len(listed[0]["responses"]) == 2


@pytest.mark.parametrize("fmt", ["csv", "json", "jsonl", "xlsx", "pdf"])
async def test_export_formats_are_pseudonymized_and_audited(
    client: httpx.AsyncClient, world: SeededWorld, login: LoginFn, catalog: dict[str, str], db: AsyncSession, fmt: str
) -> None:
    student = await login(world.student_username)
    await build_episode(client, student, catalog["B-FIG-01"])
    researcher = await login(world.researcher_email)
    response = await client.get(
        "/api/v1/research/export", params={"dataset": "interactions", "format": fmt}, headers=researcher
    )
    assert response.status_code == 200, response.text
    assert "attachment" in response.headers["content-disposition"]
    content = response.content
    if fmt == "csv":
        rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        assert rows[0]["participant_code"] == world.student_participant_code
        assert {"session_id", "event_type", "help_level", "next_student_action"} <= set(rows[0])
        assert "stu1" not in content.decode("utf-8-sig")
    elif fmt == "json":
        data = json.loads(content)
        assert data[0]["event_type"] == "SESSION_STARTED"
    elif fmt == "jsonl":
        lines = content.decode().splitlines()
        assert json.loads(lines[1])["event_type"] == "TASK_OPENED"
    elif fmt == "xlsx":
        ws = load_workbook(io.BytesIO(content)).active
        assert ws is not None and ws.cell(1, 1).value == "session_id"
    else:
        assert content.startswith(b"%PDF")
    audit = (await db.execute(select(AuditLog).where(AuditLog.action == "EXPORT_GENERATED"))).scalars().all()
    assert len(audit) == 1 and audit[0].details["format"] == fmt


async def test_export_and_participants_respect_scope(
    client: httpx.AsyncClient, world: SeededWorld, login: LoginFn, catalog: dict[str, str], db: AsyncSession
) -> None:
    import uuid

    from app.modules.identity.models import Institution
    from app.seed.demo import create_student

    institution = await db.get(Institution, uuid.UUID(world.institution_id))
    assert institution is not None
    await create_student(
        db,
        username="pending",
        password=world.password,
        institution=institution,
        grade="8",
        group_code="A",
        research_status="PENDING",
    )
    await db.commit()
    await build_episode(client, await login("pending"), catalog["B-FIG-01"])
    await build_episode(client, await login(world.student_username), catalog["B-FIG-01"])

    researcher = await login(world.researcher_email)
    participants = (await client.get("/api/v1/research/participants", headers=researcher)).json()
    assert [p["participant_code"] for p in participants] == [world.student_participant_code]
    assert set(participants[0]) == {
        "student_id",
        "participant_code",
        "grade",
        "group_code",
        "institution_code",
        "research_status",
        "sessions",
        "tasks_worked",
        "interactions",
        "episodes",
    }
    exported = json.loads(
        (
            await client.get(
                "/api/v1/research/export", params={"dataset": "interactions", "format": "json"}, headers=researcher
            )
        ).content
    )
    assert {r["participant_code"] for r in exported} == {world.student_participant_code}

    teacher = await login(world.teacher_email)
    teacher_rows = json.loads(
        (
            await client.get(
                "/api/v1/research/export", params={"dataset": "sessions", "format": "json"}, headers=teacher
            )
        ).content
    )
    assert len(teacher_rows) == 2  # el docente ve su grupo completo (incluye pendiente de consentimiento)
    denied = await client.get("/api/v1/research/export", params={"dataset": "memos", "format": "json"}, headers=teacher)
    assert denied.status_code == 403
    student = await login(world.student_username)
    assert (
        await client.get("/api/v1/research/export", params={"dataset": "interactions"}, headers=student)
    ).status_code == 403
